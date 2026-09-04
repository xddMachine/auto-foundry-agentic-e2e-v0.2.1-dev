"""Bounded, deterministic document normalisation for mission intake.

The launch boundary may receive arbitrary user documents.  This module turns
the small set of supported formats into provenance-rich excerpts while
retaining the original bytes in the read-only data room.  Extraction is best
effort: a malformed/unsupported document becomes an opaque catalog entry with
an explicit limitation.  Unsafe container structure is different and fails
closed so callers cannot accidentally inspect traversal members or unbounded
archives.

Optional dependencies are imported lazily.  PDF uses :mod:`pypdf`; XLSX uses
read-only :mod:`openpyxl`.  DOCX and ODT use constrained standard-library
ZIP/XML extraction, which keeps the core usable when those optional readers
are absent.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
import hashlib
import io
import multiprocessing
import os
from pathlib import Path, PurePosixPath
import re
import signal
import time
import zipfile
import xml.etree.ElementTree as _stdlib_ET
from typing import Any, Iterable, Mapping

try:  # Prefer the hardened parser whenever the direct dependency is present.
    from defusedxml import ElementTree as ET
except ImportError:  # pragma: no cover - exercised only in dependency-light environments
    # The package declares defusedxml as a runtime dependency.  Keep the core
    # importable before installation while applying the explicit declaration
    # rejection below before falling back to the standard parser.
    ET = _stdlib_ET


_UNSAFE_XML_DECLARATION_RE = re.compile(rb"<!\s*(?:DOCTYPE|ENTITY)\b|\b(?:SYSTEM|PUBLIC)\b", re.IGNORECASE)


DOCUMENT_INGESTION_SCHEMA_VERSION = 1
# Formats accepted into the catalog.  Only the core brief/table formats are
# normalized; the remaining known source formats are durable opaque entries so
# extraction limitations remain visible without making unrelated intake fail.
SUPPORTED_DOCUMENT_FORMATS = frozenset(
    {
        "txt",
        "text",
        "md",
        "markdown",
        "rst",
        "csv",
        "tsv",
        "pdf",
        "docx",
        "odt",
        "xlsx",
        "json",
        "jsonl",
        "ndjson",
        "parquet",
        "html",
        "htm",
        "xml",
        "log",
        "yaml",
        "yml",
        "toml",
        "ini",
        "cfg",
        "sql",
        "py",
        "sh",
    }
)
# Formats for which this module can produce bounded textual sections.  The
# broader ``SUPPORTED_DOCUMENT_FORMATS`` set is retained for standalone
# callers that want explicit opaque records for known source formats (for
# example Parquet).  A launch-time Planner catalog uses this narrower set:
# structured and binary sources already belong to the native Data Room
# catalog and must never be re-admitted as documents.
NORMALIZABLE_DOCUMENT_FORMATS = frozenset(
    {
        "txt",
        "md",
        "rst",
        "csv",
        "tsv",
        "pdf",
        "docx",
        "odt",
        "xlsx",
    }
)
# Intake does not impose a business-size or document-count ceiling by
# default.  Structural archive checks (safe paths, duplicate names, supported
# compression, CRC/read verification, and the compression-ratio bomb guard)
# remain mandatory.  Callers may provide explicit content/count caps; those
# are represented as limited/opaque records whenever extraction is skipped.
DEFAULT_MAX_DOCUMENT_BYTES: int | None = None
DEFAULT_MAX_MEMBER_BYTES: int | None = None
DEFAULT_MAX_TOTAL_BYTES: int | None = None
DEFAULT_MAX_ARCHIVE_ENTRIES: int | None = None
DEFAULT_MAX_EXCERPT_BYTES = 16 * 1024
DEFAULT_MAX_EXCERPTS = 128
DEFAULT_MAX_ROWS = 128
DEFAULT_MAX_COLUMNS = 128
MAX_COMPRESSION_RATIO = 1000.0
DEFAULT_MAX_PDF_PAGES = 128
DEFAULT_MAX_PDF_OUTPUT_BYTES = 16 * 1024 * 1024
DEFAULT_PDF_TIMEOUT_SECONDS = 8.0
DEFAULT_PDF_CPU_SECONDS = 8
DEFAULT_PDF_MEMORY_BYTES = 256 * 1024 * 1024
# Catalog-wide limits prevent a large safe data room from holding the launch
# lock for one per-file timeout per member.  Per-document limits remain useful
# for hostile individual files; these aggregate limits are the final admission
# budget for one catalog pass.
DEFAULT_MAX_PARSED_PDFS = 32
# Public alias for callers that describe the same bound in terms of spawned
# parser workers rather than successfully parsed files.
DEFAULT_MAX_PDF_PROCESSES = DEFAULT_MAX_PARSED_PDFS
DEFAULT_MAX_PDF_TOTAL_WALL_SECONDS = 30.0
DEFAULT_MAX_PDF_TOTAL_OUTPUT_BYTES = 32 * 1024 * 1024
DEFAULT_MAX_TOTAL_NORMALIZED_TEXT_BYTES = 64 * 1024 * 1024

_REVISION_SUFFIX_RE = re.compile(r"~sha256-[0-9a-f]{64}(?=(?:\.[^/]*)?$)", re.IGNORECASE)


class DocumentIngestionError(ValueError):
    """Base error for deterministic ingestion failures."""


class UnsafeDocumentArchiveError(DocumentIngestionError):
    """Archive structure is unsafe and must fail admission."""


class _ArchiveMemberReadError(zipfile.BadZipFile):
    """A member failed while streaming, retaining any verified bytes hash."""

    def __init__(self, message: str, *, observed_hash: str | None, observed_size: int) -> None:
        super().__init__(message)
        self.observed_hash = observed_hash
        self.observed_size = observed_size


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _file_sha256(path: Path) -> tuple[str, int]:
    digest = hashlib.sha256()
    size = 0
    with path.open("rb") as stream:
        while chunk := stream.read(1024 * 1024):
            digest.update(chunk)
            size += len(chunk)
    return digest.hexdigest(), size


def _safe_member_name(name: str) -> str:
    raw = str(name).replace("\\", "/")
    if not raw or raw.startswith("/") or raw.startswith("~") or "\x00" in raw:
        raise UnsafeDocumentArchiveError("archive member path is not relative")
    parts = raw.split("/")
    if any(part in {"", ".", ".."} for part in parts):
        raise UnsafeDocumentArchiveError("archive member path contains traversal")
    if len(parts[0]) == 2 and parts[0][1] == ":":
        raise UnsafeDocumentArchiveError("archive member path is absolute")
    return "/".join(parts)


def _is_symlink(info: zipfile.ZipInfo) -> bool:
    mode = (info.external_attr >> 16) & 0xFFFF
    return (mode & 0o170000) == 0o120000


def _format(path_or_name: str | Path, explicit: str | None = None) -> str:
    value = (explicit or Path(path_or_name).suffix.lstrip(".") or "binary").lower().lstrip(".")
    if value == "text":
        return "txt"
    if value in {"markdown", "mdown"}:
        return "md"
    if value == "tsv":
        return "tsv"
    return value


def _source_path_for_ref(document_ref: str) -> str:
    """Return the display/source path represented by a revision-qualified ref."""

    value = str(document_ref).replace("\\", "/")
    return _REVISION_SUFFIX_RE.sub("", value)


def revision_document_ref(source_path: str | Path, content_hash: str) -> str:
    """Build a deterministic revision ref while retaining the file suffix.

    A suffix-preserving ref (``brief~sha256-<digest>.md``) keeps format
    dispatch deterministic while making changed bytes impossible to overwrite
    an older catalog entry.  ``source_path`` is display metadata and may be
    reused across revisions.
    """

    source = str(source_path).replace("\\", "/")
    digest = str(content_hash).lower()
    if not re.fullmatch(r"[0-9a-f]{64}", digest):
        raise ValueError("content_hash must be a SHA-256 hex digest")
    if _REVISION_SUFFIX_RE.search(source):
        return source
    path = PurePosixPath(source)
    suffix = path.suffix
    stem = path.name[: -len(suffix)] if suffix else path.name
    name = f"{stem}~sha256-{digest}{suffix}"
    parent = str(path.parent)
    return name if parent in {"", "."} else f"{parent}/{name}"


def _decode(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="replace")


def _truncate(text: str, cap: int) -> tuple[str, bool]:
    encoded = text.encode("utf-8")
    if len(encoded) <= cap:
        return text, False
    # Decode a bounded prefix without splitting a UTF-8 sequence.
    result = encoded[:cap].decode("utf-8", errors="ignore")
    return result, True


def _xml_text(element: ET.Element) -> str:
    return " ".join(part.strip() for part in element.itertext() if part and part.strip())


@dataclass(frozen=True)
class NormalizedDocumentSection:
    document_ref: str
    format: str
    locator: Mapping[str, Any]
    text: str
    content_hash: str
    limitations: tuple[str, ...] = ()
    kind: str = "section"

    def __post_init__(self) -> None:
        if not str(self.document_ref).strip():
            raise ValueError("document_ref must not be empty")
        fmt = _format(self.document_ref, self.format)
        if fmt not in SUPPORTED_DOCUMENT_FORMATS:
            raise ValueError(f"unsupported document format: {fmt}")
        object.__setattr__(self, "format", fmt)
        if not isinstance(self.locator, Mapping):
            raise TypeError("locator must be an object")
        object.__setattr__(self, "locator", {str(k): v for k, v in self.locator.items()})
        if not isinstance(self.text, str):
            raise TypeError("section text must be a string")
        digest = str(self.content_hash).lower()
        if not re.fullmatch(r"[0-9a-f]{64}", digest):
            raise ValueError("section content_hash must be a SHA-256 digest")
        if digest != _sha256(self.text.encode("utf-8")):
            raise ValueError("section content_hash does not match text")
        object.__setattr__(self, "content_hash", digest)
        object.__setattr__(self, "limitations", tuple(str(value) for value in (self.limitations or ())))
        object.__setattr__(self, "kind", str(self.kind or "section"))

    @property
    def locator_ref(self) -> Mapping[str, Any]:
        return self.locator

    @property
    def documentRef(self) -> str:  # pragma: no cover - wire compatibility alias
        return self.document_ref

    @property
    def contentHash(self) -> str:  # pragma: no cover - wire compatibility alias
        return self.content_hash

    def __getitem__(self, key: str) -> Any:
        if key in {"document_ref", "documentRef"}:
            return self.document_ref
        if key == "format":
            return self.format
        if key in {"content_hash", "contentHash"}:
            return self.content_hash
        if key == "text":
            return self.text
        if key in {"locator", "location"}:
            return self.locator
        if key in self.locator:
            return self.locator[key]
        if key == "limitations":
            return self.limitations
        raise KeyError(key)

    def to_dict(self) -> dict[str, Any]:
        return {
            "document_ref": self.document_ref,
            "format": self.format,
            "locator": dict(self.locator),
            "text": self.text,
            "content_hash": self.content_hash,
            "limitations": list(self.limitations),
            "kind": self.kind,
        }

    @classmethod
    def from_dict(cls, value: Mapping[str, Any] | "NormalizedDocumentSection") -> "NormalizedDocumentSection":
        if isinstance(value, cls):
            return value
        if not isinstance(value, Mapping):
            raise TypeError("document section must be an object")
        raw = dict(value)
        raw["document_ref"] = raw.get("document_ref", raw.get("documentRef"))
        raw.setdefault("format", _format(raw["document_ref"]))
        raw["content_hash"] = raw.get("content_hash", raw.get("contentHash"))
        raw["locator"] = raw.get("locator", raw.get("location", {}))
        raw.setdefault("limitations", ())
        return cls(**raw)


@dataclass(frozen=True)
class NormalizedDocument:
    document_ref: str
    format: str
    content_hash: str
    size_bytes: int
    sections: tuple[NormalizedDocumentSection, ...] = ()
    limitations: tuple[str, ...] = ()
    extraction: str = "normalized"
    # ``document_ref`` is the immutable catalog identity.  ``source_path`` is
    # intentionally display-only and can recur when a later data-room
    # revision contains changed bytes for the same path.
    source_path: str | None = None

    def __post_init__(self) -> None:
        if not str(self.document_ref).strip():
            raise ValueError("document_ref must not be empty")
        object.__setattr__(self, "format", _format(self.document_ref, self.format))
        source_path = str(self.source_path or _source_path_for_ref(self.document_ref)).replace("\\", "/")
        if not source_path.strip():
            raise ValueError("source_path must not be empty")
        object.__setattr__(self, "source_path", source_path)
        digest = str(self.content_hash).lower()
        if not re.fullmatch(r"[0-9a-f]{64}", digest):
            raise ValueError("content_hash must be a SHA-256 digest")
        object.__setattr__(self, "content_hash", digest)
        if isinstance(self.size_bytes, bool) or not isinstance(self.size_bytes, int) or self.size_bytes < 0:
            raise ValueError("size_bytes must be a non-negative integer")
        sections = tuple(
            section if isinstance(section, NormalizedDocumentSection) else NormalizedDocumentSection(**dict(section))
            for section in (self.sections or ())
        )
        # Whitespace-only excerpts carry no usable evidence and must not be
        # exposed as document bindings.  The host remains the sole authority
        # for canonical text; dropping them here also covers catalog values
        # reconstructed from wire dictionaries.
        sections = tuple(section for section in sections if section.text.strip())
        if any(section.document_ref != self.document_ref or section.format != self.format for section in sections):
            raise ValueError("document sections must bind to their parent document")
        object.__setattr__(self, "sections", sections)
        object.__setattr__(self, "limitations", tuple(str(value) for value in (self.limitations or ())))
        extraction = str(self.extraction or "normalized")
        if extraction not in {"normalized", "opaque", "limited"}:
            raise ValueError("unsupported extraction status")
        object.__setattr__(self, "extraction", extraction)

    @property
    def document_hash(self) -> str:
        return self.content_hash

    @property
    def documentRef(self) -> str:  # pragma: no cover - wire compatibility alias
        return self.document_ref

    @property
    def contentHash(self) -> str:  # pragma: no cover - wire compatibility alias
        return self.content_hash

    @property
    def sourcePath(self) -> str:  # pragma: no cover - wire compatibility alias
        return str(self.source_path)

    def __getitem__(self, key: str) -> Any:
        aliases = {
            "document_ref": "document_ref",
            "documentRef": "document_ref",
            "format": "format",
            "content_hash": "content_hash",
            "contentHash": "content_hash",
            "size_bytes": "size_bytes",
            "sizeBytes": "size_bytes",
            "sections": "sections",
            "limitations": "limitations",
            "extraction": "extraction",
            "source_path": "source_path",
            "sourcePath": "source_path",
        }
        if key not in aliases:
            raise KeyError(key)
        return getattr(self, aliases[key])

    def to_dict(self) -> dict[str, Any]:
        return {
            "document_ref": self.document_ref,
            "format": self.format,
            "content_hash": self.content_hash,
            "size_bytes": self.size_bytes,
            "sections": [section.to_dict() for section in self.sections],
            "limitations": list(self.limitations),
            "extraction": self.extraction,
            "source_path": self.source_path,
        }

    @classmethod
    def from_dict(cls, value: Mapping[str, Any] | "NormalizedDocument") -> "NormalizedDocument":
        if isinstance(value, cls):
            return value
        if not isinstance(value, Mapping):
            raise TypeError("normalized document must be an object")
        raw = dict(value)
        raw["document_ref"] = raw.get("document_ref", raw.get("documentRef"))
        raw.pop("documentRef", None)
        raw.setdefault("format", _format(raw["document_ref"]))
        raw["content_hash"] = raw.get("content_hash", raw.get("contentHash"))
        raw.pop("contentHash", None)
        raw["size_bytes"] = raw.get("size_bytes", raw.get("sizeBytes"))
        raw.pop("sizeBytes", None)
        raw.setdefault("sections", ())
        raw["sections"] = tuple(NormalizedDocumentSection.from_dict(item) for item in raw["sections"])
        raw.setdefault("limitations", ())
        raw.setdefault("extraction", "normalized")
        raw["source_path"] = raw.get("source_path", raw.get("sourcePath"))
        raw.pop("sourcePath", None)
        return cls(**raw)


@dataclass(frozen=True)
class DocumentCatalog:
    documents: tuple[NormalizedDocument, ...] = ()
    schema_version: int = DOCUMENT_INGESTION_SCHEMA_VERSION
    limitations: tuple[str, ...] = ()

    def __post_init__(self) -> None:
        if isinstance(self.schema_version, bool) or self.schema_version != DOCUMENT_INGESTION_SCHEMA_VERSION:
            raise ValueError("unsupported document catalog schema version")
        documents = tuple(
            document if isinstance(document, NormalizedDocument) else NormalizedDocument.from_dict(document)
            for document in (self.documents or ())
        )
        refs = [document.document_ref for document in documents]
        if len(refs) != len(set(refs)):
            raise ValueError("document catalog references must be unique")
        object.__setattr__(self, "documents", documents)
        object.__setattr__(self, "limitations", tuple(str(value) for value in (self.limitations or ())))

    def __iter__(self):
        return iter(self.documents)

    def __len__(self) -> int:
        return len(self.documents)

    def __getitem__(self, index):
        return self.documents[index]

    @property
    def entries(self) -> tuple[NormalizedDocument, ...]:
        return self.documents

    def to_dict(self) -> dict[str, Any]:
        return {
            "schema_version": self.schema_version,
            "documents": [document.to_dict() for document in self.documents],
            "limitations": list(self.limitations),
        }

    @classmethod
    def from_dict(cls, value: Mapping[str, Any] | "DocumentCatalog") -> "DocumentCatalog":
        if isinstance(value, cls):
            return value
        if not isinstance(value, Mapping):
            raise TypeError("document catalog must be an object")
        raw_documents = value.get("documents", value.get("entries", ()))
        if not isinstance(raw_documents, (list, tuple)):
            raise TypeError("document catalog documents must be a list")
        documents: list[NormalizedDocument] = []
        for raw in raw_documents:
            if isinstance(raw, NormalizedDocument):
                documents.append(raw)
                continue
            if not isinstance(raw, Mapping):
                raise TypeError("document catalog entry must be an object")
            item = dict(raw)
            item["document_ref"] = item.get("document_ref", item.get("documentRef"))
            item["content_hash"] = item.get("content_hash", item.get("contentHash"))
            item["size_bytes"] = item.get("size_bytes", item.get("sizeBytes"))
            item["source_path"] = item.get("source_path", item.get("sourcePath"))
            item.pop("documentRef", None)
            item.pop("contentHash", None)
            item.pop("sizeBytes", None)
            item.pop("sourcePath", None)
            sections = item.get("sections", ())
            normalized_sections: list[NormalizedDocumentSection] = []
            for section in sections or ():
                if isinstance(section, NormalizedDocumentSection):
                    normalized_sections.append(section)
                else:
                    section = dict(section)
                    section["document_ref"] = section.get("document_ref", section.get("documentRef", item["document_ref"]))
                    section["content_hash"] = section.get("content_hash", section.get("contentHash"))
                    normalized_sections.append(NormalizedDocumentSection.from_dict(section))
            item["sections"] = tuple(normalized_sections)
            documents.append(NormalizedDocument(**item))
        schema_version = value.get("schema_version", value.get("schemaVersion", DOCUMENT_INGESTION_SCHEMA_VERSION))
        return cls(documents=tuple(documents), schema_version=schema_version, limitations=value.get("limitations", ()))

    def planner_payload(
        self,
        *,
        max_excerpt_bytes: int = DEFAULT_MAX_EXCERPT_BYTES,
        max_excerpts: int = DEFAULT_MAX_EXCERPTS,
    ) -> dict[str, Any]:
        """Return deterministic, bounded catalog data suitable for a Planner."""

        if max_excerpt_bytes < 0 or max_excerpts < 0:
            raise ValueError("planner catalog limits cannot be negative")
        used = 0
        emitted = 0
        documents: list[dict[str, Any]] = []
        for document in sorted(self.documents, key=lambda item: item.document_ref):
            sections: list[dict[str, Any]] = []
            for section in sorted(document.sections, key=lambda item: repr(sorted(item.locator.items()))):
                if not section.text.strip():
                    continue
                if emitted >= max_excerpts:
                    break
                remaining = max_excerpt_bytes - used
                if remaining <= 0:
                    break
                excerpt, truncated = _truncate(section.text, remaining)
                if not excerpt and section.text:
                    break
                section_payload = section.to_dict()
                section_payload["text"] = excerpt
                if truncated:
                    section_payload["limitations"] = [*section_payload["limitations"], "planner excerpt byte cap reached"]
                encoded_size = len(excerpt.encode("utf-8"))
                used += encoded_size
                emitted += 1
                sections.append(section_payload)
            documents.append({
                "document_ref": document.document_ref,
                "source_path": document.source_path,
                "format": document.format,
                "content_hash": document.content_hash,
                "size_bytes": document.size_bytes,
                "extraction": document.extraction,
                "limitations": list(document.limitations),
                "sections": sections,
            })
        return {
            "schema_version": self.schema_version,
            "documents": documents,
            "limitations": list(self.limitations),
            "bounded": True,
            "max_excerpt_bytes": max_excerpt_bytes,
            "max_excerpts": max_excerpts,
        }


def _rebind_document(document: NormalizedDocument, document_ref: str) -> NormalizedDocument:
    """Copy one immutable document under a new revision-qualified ref."""

    if document.document_ref == document_ref:
        return document
    sections = tuple(
        NormalizedDocumentSection(
            document_ref=document_ref,
            format=section.format,
            locator=section.locator,
            text=section.text,
            content_hash=section.content_hash,
            limitations=section.limitations,
            kind=section.kind,
        )
        for section in document.sections
    )
    return NormalizedDocument(
        document_ref=document_ref,
        format=document.format,
        content_hash=document.content_hash,
        size_bytes=document.size_bytes,
        sections=sections,
        limitations=document.limitations,
        extraction=document.extraction,
        source_path=document.source_path,
    )


def revision_qualify_catalog(
    parent: DocumentCatalog | Mapping[str, Any] | None,
    child: DocumentCatalog | Mapping[str, Any] | None,
) -> tuple[DocumentCatalog | None, dict[str, str]]:
    """Bind child entries to immutable refs relative to a parent catalog.

    The returned map translates child ``document_ref`` values to the refs that
    are valid in the cumulative catalog.  A same-path/same-hash child reuses
    the existing immutable entry; changed bytes receive a deterministic
    ``~sha256-<digest>`` suffix while retaining ``source_path`` metadata.
    Catalog order is preserved for newly admitted entries and parent entries
    are never replaced.
    """

    parent_catalog = None if parent is None else DocumentCatalog.from_dict(parent)
    child_catalog = None if child is None else DocumentCatalog.from_dict(child)
    if child_catalog is None:
        return parent_catalog, {}
    if parent_catalog is None:
        return child_catalog, {document.document_ref: document.document_ref for document in child_catalog.documents}
    parent_documents = tuple(parent_catalog.documents)
    parent_by_path: dict[str, list[NormalizedDocument]] = {}
    for document in parent_documents:
        path = str(document.source_path or _source_path_for_ref(document.document_ref))
        parent_by_path.setdefault(path, []).append(document)
    used_refs = {document.document_ref for document in parent_documents}
    merged = list(parent_documents)
    translations: dict[str, str] = {}
    for document in child_catalog.documents:
        source_path = str(document.source_path or _source_path_for_ref(document.document_ref))
        candidates = parent_by_path.get(source_path, ())
        same_hash = next(
            (candidate for candidate in sorted(candidates, key=lambda value: value.document_ref) if candidate.content_hash == document.content_hash),
            None,
        )
        if same_hash is not None:
            translations[document.document_ref] = same_hash.document_ref
            continue
        target_ref = document.document_ref
        if candidates or target_ref in used_refs:
            target_ref = revision_document_ref(source_path, document.content_hash)
        # A pathological pre-qualified ref collision should never overwrite an
        # unrelated source.  The digest makes this deterministic; retain a
        # qualified ref if a caller supplied one already.
        if target_ref in used_refs:
            existing = next((value for value in merged if value.document_ref == target_ref), None)
            if existing is not None and existing.content_hash == document.content_hash and str(existing.source_path) == source_path:
                translations[document.document_ref] = target_ref
                continue
            target_ref = revision_document_ref(source_path, document.content_hash)
        rebound = _rebind_document(document, target_ref)
        translations[document.document_ref] = target_ref
        merged.append(rebound)
        used_refs.add(target_ref)
        parent_by_path.setdefault(source_path, []).append(rebound)
    limitations = tuple(dict.fromkeys((*parent_catalog.limitations, *child_catalog.limitations)))
    return DocumentCatalog(tuple(merged), limitations=limitations), translations


def _section(document_ref: str, fmt: str, locator: Mapping[str, Any], text: str, *, limitations: Iterable[str] = (), kind: str = "section") -> NormalizedDocumentSection:
    return NormalizedDocumentSection(
        document_ref=document_ref,
        format=fmt,
        locator=locator,
        text=text,
        content_hash=_sha256(text.encode("utf-8")),
        limitations=tuple(limitations),
        kind=kind,
    )


def _normalize_text(document_ref: str, fmt: str, data: bytes, *, max_excerpt_bytes: int) -> tuple[tuple[NormalizedDocumentSection, ...], tuple[str, ...], str]:
    text = _decode(data)
    limitations: list[str] = []
    excerpt, truncated = _truncate(text, max_excerpt_bytes)
    if truncated:
        limitations.append("document excerpt truncated at configured byte cap")
    # Keep line/paragraph provenance while avoiding one enormous section.  A
    # plain document with no newline still gets one stable section.
    raw_parts = text.splitlines() or ([text] if text else [])
    sections: list[NormalizedDocumentSection] = []
    consumed = 0
    for index, part in enumerate(raw_parts, start=1):
        if consumed >= max_excerpt_bytes:
            break
        available = max_excerpt_bytes - consumed
        value, was_truncated = _truncate(part, available)
        if not value and part:
            break
        section_limits = [*limitations] if was_truncated else []
        sections.append(_section(document_ref, fmt, {"section": index, "paragraph": index}, value, limitations=section_limits, kind="section"))
        consumed += len(value.encode("utf-8"))
        if was_truncated:
            break
    if not sections and excerpt:
        sections.append(_section(document_ref, fmt, {"section": 1, "paragraph": 1}, excerpt, limitations=limitations))
    return tuple(sections), tuple(limitations), "limited" if limitations else "normalized"


def _normalize_delimited(document_ref: str, fmt: str, data: bytes, *, max_excerpt_bytes: int, max_rows: int, max_columns: int) -> tuple[tuple[NormalizedDocumentSection, ...], tuple[str, ...], str]:
    text = _decode(data)
    delimiter = "\t" if fmt == "tsv" else ","
    limitations: list[str] = []
    sections: list[NormalizedDocumentSection] = []
    try:
        rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    except (csv.Error, UnicodeError) as exc:
        return (), (f"tabular extraction failed: {exc}",), "opaque"
    if len(rows) > max_rows:
        limitations.append("table rows truncated at configured cap")
    header = rows[0] if rows else []
    for row_number, row in enumerate(rows[:max_rows], start=1):
        values = row[:max_columns]
        if len(row) > max_columns:
            limitations.append("table columns truncated at configured cap")
        cells = [f"{header[index] if index < len(header) and header[index] else 'column_' + str(index + 1)}={value}" for index, value in enumerate(values)]
        row_text = " | ".join(cells)
        remaining = max_excerpt_bytes - sum(len(value.text.encode("utf-8")) for value in sections)
        if remaining <= 0:
            limitations.append("table excerpt truncated at configured byte cap")
            break
        row_text, truncated = _truncate(row_text, remaining)
        if truncated:
            limitations.append("table excerpt truncated at configured byte cap")
        cell_refs = [f"{_column_name(index + 1)}{row_number}" for index in range(len(values))]
        sections.append(_section(document_ref, fmt, {"section": 1, "row": row_number, "cell": cell_refs[0] if cell_refs else None, "cells": cell_refs}, row_text, limitations=(["row excerpt truncated"] if truncated else ()), kind="table_excerpt"))
        if truncated:
            break
    return tuple(sections), tuple(dict.fromkeys(limitations)), "limited" if limitations else "normalized"


def _safe_xml_fromstring(data: bytes) -> Any:
    """Parse one bounded XML member without DTD/entity processing.

    ``defusedxml`` is the normal parser and rejects XXE, entity expansion,
    and other hostile XML constructs.  A dependency-light source checkout can
    still import the module before installation; in that narrow fallback we
    reject every DTD/entity/external-identifier declaration before invoking
    ``xml.etree``.  This keeps the fallback fail-closed rather than silently
    enabling an unsafe parser.
    """

    if _UNSAFE_XML_DECLARATION_RE.search(data):
        raise DocumentIngestionError("XML DTD/entity declarations are not permitted")
    try:
        return ET.fromstring(data)
    except (ET.ParseError, _stdlib_ET.ParseError) as exc:
        raise DocumentIngestionError(f"XML parsing failed: {exc}") from exc


def _column_name(index: int) -> str:
    value = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        value = chr(65 + remainder) + value
    return value


def _validate_archive(
    data: bytes | bytearray | memoryview | str | Path,
    *,
    max_member_bytes: int | None,
    max_total_bytes: int | None,
    max_entries: int | None,
    enforce_resource_limits: bool = True,
) -> tuple[tuple[zipfile.ZipInfo, ...], int]:
    # Keep path-backed archives on disk.  ``ZipFile`` reads only the central
    # directory and each member is streamed by the caller; converting a large
    # data-room archive into one ``bytes`` object would defeat the bounded
    # admission boundary.
    archive_source: Any
    if isinstance(data, (str, Path)):
        archive_source = data
    else:
        archive_source = io.BytesIO(bytes(data))
    try:
        archive = zipfile.ZipFile(archive_source, "r")
    except (OSError, zipfile.BadZipFile, RuntimeError) as exc:
        raise UnsafeDocumentArchiveError(f"invalid document archive: {exc}") from exc
    with archive:
        infos = archive.infolist()
        if enforce_resource_limits and max_entries is not None and len(infos) > max_entries:
            raise UnsafeDocumentArchiveError("document archive has too many members")
        total = 0
        seen: set[str] = set()
        for info in infos:
            name = _safe_member_name(info.filename)
            if info.is_dir():
                continue
            if _is_symlink(info) or info.flag_bits & 0x1:
                raise UnsafeDocumentArchiveError(f"unsafe document archive member: {name}")
            if info.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED, zipfile.ZIP_BZIP2, zipfile.ZIP_LZMA}:
                raise UnsafeDocumentArchiveError(f"unsupported document archive compression: {name}")
            if enforce_resource_limits and max_member_bytes is not None and info.file_size > max_member_bytes:
                raise UnsafeDocumentArchiveError(f"document archive member exceeds size cap: {name}")
            # Compression-ratio validation is structural protection against
            # decompression bombs, not a business-size cap.  It therefore stays
            # active even when optional resource caps are soft.
            if info.file_size < 0 or info.compress_size < 0:
                raise UnsafeDocumentArchiveError(f"invalid document archive member size: {name}")
            if info.file_size and info.compress_size == 0:
                raise UnsafeDocumentArchiveError(f"document archive compression ratio is invalid: {name}")
            if info.compress_size and info.file_size / info.compress_size > MAX_COMPRESSION_RATIO:
                raise UnsafeDocumentArchiveError(f"document archive compression ratio exceeds cap: {name}")
            total += info.file_size
            if enforce_resource_limits and max_total_bytes is not None and total > max_total_bytes:
                raise UnsafeDocumentArchiveError("document archive exceeds expanded size cap")
            key = name.casefold()
            if key in seen:
                raise UnsafeDocumentArchiveError(f"duplicate document archive member: {name}")
            seen.add(key)
        # ZipInfo values remain valid after this pass; the caller reopens the
        # bytes for actual reads to keep the validation/read boundary clear.
        return tuple(infos), total


def _normalize_docx_or_odt(document_ref: str, fmt: str, data: bytes, *, max_member_bytes: int | None, max_total_bytes: int | None, max_excerpt_bytes: int, max_entries: int | None) -> tuple[tuple[NormalizedDocumentSection, ...], tuple[str, ...], str]:
    infos, _total = _validate_archive(data, max_member_bytes=max_member_bytes, max_total_bytes=max_total_bytes, max_entries=max_entries)
    wanted = "word/document.xml" if fmt == "docx" else "content.xml"
    try:
        with zipfile.ZipFile(io.BytesIO(data), "r") as archive:
            candidates = [info for info in infos if info.filename == wanted]
            if not candidates:
                raise DocumentIngestionError(f"{wanted} is missing")
            payload = archive.read(candidates[0])
        if max_member_bytes is not None and len(payload) > max_member_bytes:
            raise DocumentIngestionError("document XML exceeds extraction cap")
        root = _safe_xml_fromstring(payload)
    except (OSError, zipfile.BadZipFile, UnicodeError, KeyError, DocumentIngestionError) as exc:
        return (), (f"{fmt.upper()} text extraction failed: {exc}",), "opaque"
    sections: list[NormalizedDocumentSection] = []
    if fmt == "docx":
        # Paragraph and table-cell elements carry stable ordinal locators.
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        paragraphs = root.iter(namespace + "p")
        for index, paragraph in enumerate(paragraphs, start=1):
            value = _xml_text(paragraph)
            if value:
                sections.append(_section(document_ref, fmt, {"paragraph": index}, value, kind="paragraph"))
    else:
        paragraphs = []
        for element in root.iter():
            local = element.tag.rsplit("}", 1)[-1]
            if local in {"p", "h"}:
                value = _xml_text(element)
                if value:
                    paragraphs.append(value)
        for index, value in enumerate(paragraphs, start=1):
            sections.append(_section(document_ref, fmt, {"paragraph": index}, value, kind="paragraph"))
    used = 0
    bounded: list[NormalizedDocumentSection] = []
    for section in sections:
        if used >= max_excerpt_bytes:
            break
        value, truncated = _truncate(section.text, max_excerpt_bytes - used)
        if not value and section.text:
            break
        bounded.append(_section(document_ref, fmt, section.locator, value, limitations=("document excerpt truncated at configured byte cap",) if truncated else (), kind=section.kind))
        used += len(value.encode("utf-8"))
        if truncated:
            break
    limitations = ("document excerpt truncated at configured byte cap",) if len(bounded) < len(sections) else ()
    return tuple(bounded), limitations, "limited" if limitations else "normalized"


def _pdf_extract_worker(
    data: bytes,
    *,
    max_excerpt_bytes: int,
    max_pages: int,
    max_output_bytes: int,
    cpu_seconds: int,
    memory_bytes: int,
    connection: Any,
) -> None:
    """Extract PDF text in an isolated process with OS resource bounds."""

    try:
        # ``resource`` is Unix-only; the parent wall-clock bound remains in
        # force on platforms where these optional limits are unavailable.
        import resource

        if cpu_seconds > 0:
            resource.setrlimit(resource.RLIMIT_CPU, (cpu_seconds, cpu_seconds + 1))
        if memory_bytes > 0 and hasattr(resource, "RLIMIT_AS"):
            resource.setrlimit(resource.RLIMIT_AS, (memory_bytes, memory_bytes))
    except (ImportError, OSError, ValueError):
        pass
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data), strict=False)
        page_count = len(reader.pages)
        if page_count > max_pages:
            connection.send(
                {
                    "status": "limited",
                    "sections": [],
                    "limitations": [f"PDF page count exceeds configured cap ({page_count}>{max_pages})"],
                }
            )
            return
        sections: list[dict[str, Any]] = []
        limitations: list[str] = []
        used = 0
        for page_number in range(1, page_count + 1):
            page = reader.pages[page_number - 1]
            try:
                value = str(page.extract_text() or "").strip()
            except Exception as exc:
                limitations.append(f"PDF page {page_number} extraction failed: {exc.__class__.__name__}")
                continue
            if not value:
                continue
            remaining = min(max_excerpt_bytes - used, max_output_bytes - used)
            if remaining <= 0:
                limitations.append("PDF extraction output exceeds configured byte cap")
                break
            excerpt, truncated = _truncate(value, remaining)
            if not excerpt:
                limitations.append("PDF extraction output exceeds configured byte cap")
                break
            sections.append(
                {
                    "page": page_number,
                    "text": excerpt,
                    "truncated": truncated,
                }
            )
            used += len(excerpt.encode("utf-8"))
            if truncated:
                limitations.append("PDF excerpts truncated at configured byte cap")
                if used >= max_output_bytes:
                    limitations.append("PDF extraction output exceeds configured byte cap")
                break
        connection.send({"status": "limited" if limitations else "normalized", "sections": sections, "limitations": limitations})
    except Exception as exc:  # pypdf exposes several parser-specific errors
        try:
            connection.send(
                {
                    "status": "opaque",
                    "sections": [],
                    "limitations": [f"PDF extraction failed: {exc.__class__.__name__}: {exc}"],
                }
            )
        except (BrokenPipeError, OSError):
            pass
    finally:
        try:
            connection.close()
        except (BrokenPipeError, OSError):
            pass


def _normalize_pdf(
    document_ref: str,
    data: bytes,
    *,
    max_excerpt_bytes: int,
    max_pages: int = DEFAULT_MAX_PDF_PAGES,
    max_output_bytes: int = DEFAULT_MAX_PDF_OUTPUT_BYTES,
    timeout_seconds: float = DEFAULT_PDF_TIMEOUT_SECONDS,
    cpu_seconds: int = DEFAULT_PDF_CPU_SECONDS,
    memory_bytes: int = DEFAULT_PDF_MEMORY_BYTES,
) -> tuple[tuple[NormalizedDocumentSection, ...], tuple[str, ...], str]:
    try:
        from pypdf import PdfReader as _PdfReader  # noqa: F401 - availability probe
    except ImportError:
        return (), ("PDF extraction unavailable: optional pypdf dependency is not installed",), "opaque"
    if max_pages <= 0 or max_output_bytes <= 0 or timeout_seconds <= 0:
        return (), ("PDF extraction limits are invalid",), "opaque"
    # ``fork`` is unsafe when the Control Center HTTP server has worker
    # threads (and can inherit locked pypdf/parser state).  Always use a
    # fresh interpreter; the bounded wall-clock/rlimit guard handles workers
    # that fail to initialise as a limited extraction.
    try:
        context = multiprocessing.get_context("spawn")
    except (ValueError, RuntimeError):
        context = multiprocessing.get_context()
    parent_connection, child_connection = context.Pipe(duplex=False)
    process = context.Process(
        target=_pdf_extract_worker,
        kwargs={
            "data": bytes(data),
            "max_excerpt_bytes": int(max_excerpt_bytes),
            "max_pages": int(max_pages),
            "max_output_bytes": int(max_output_bytes),
            "cpu_seconds": int(cpu_seconds),
            "memory_bytes": int(memory_bytes),
            "connection": child_connection,
        },
    )
    try:
        process.start()
        child_connection.close()
        process.join(float(timeout_seconds))
        if process.is_alive():
            process.terminate()
            process.join(0.25)
            if process.is_alive():
                # A worker can ignore SIGTERM (or a test double can model a
                # stuck parser).  Kill explicitly, then perform a bounded
                # reap so the launch lock cannot retain a zombie child.
                killer = getattr(process, "kill", None)
                if callable(killer):
                    killer()
                elif getattr(process, "pid", None):
                    try:
                        os.kill(int(process.pid), signal.SIGKILL)
                    except (OSError, TypeError, ValueError):
                        pass
                process.join(1.0)
            return (), ("PDF extraction exceeded configured wall-time limit",), "limited"
        if not parent_connection.poll(0.1):
            return (), ("PDF extraction failed before returning a bounded result",), "limited"
        result = parent_connection.recv()
    except (OSError, EOFError, ValueError) as exc:
        return (), (f"PDF extraction process failed: {exc.__class__.__name__}: {exc}",), "limited"
    finally:
        try:
            close_process = getattr(process, "close", None)
            if callable(close_process) and not process.is_alive():
                close_process()
        except (OSError, AttributeError, ValueError):
            pass
        try:
            child_connection.close()
        except (BrokenPipeError, OSError, AttributeError):
            pass
        try:
            parent_connection.close()
        except (OSError, AttributeError):
            pass
    if not isinstance(result, Mapping):
        return (), ("PDF extraction returned an invalid bounded result",), "limited"
    limitations = [str(value) for value in result.get("limitations", ())]
    sections: list[NormalizedDocumentSection] = []
    for raw in result.get("sections", ()):
        if not isinstance(raw, Mapping) or not isinstance(raw.get("page"), int) or not isinstance(raw.get("text"), str):
            limitations.append("PDF extraction returned an invalid section")
            continue
        page_limits = ("PDF excerpts truncated at configured byte cap",) if raw.get("truncated") else ()
        sections.append(_section(document_ref, "pdf", {"page": raw["page"]}, raw["text"], limitations=page_limits, kind="page"))
    status = str(result.get("status") or "limited")
    if status not in {"normalized", "limited", "opaque"}:
        status = "limited"
    return tuple(sections), tuple(dict.fromkeys(limitations)), status


def _normalize_xlsx(document_ref: str, data: bytes, *, max_excerpt_bytes: int, max_rows: int, max_columns: int, max_member_bytes: int | None, max_total_bytes: int | None, max_entries: int | None) -> tuple[tuple[NormalizedDocumentSection, ...], tuple[str, ...], str]:
    try:
        _validate_archive(data, max_member_bytes=max_member_bytes, max_total_bytes=max_total_bytes, max_entries=max_entries)
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except ImportError:
        return (), ("XLSX extraction unavailable: optional openpyxl dependency is not installed",), "opaque"
    except Exception as exc:
        return (), (f"XLSX extraction failed: {exc.__class__.__name__}: {exc}",), "opaque"
    sections: list[NormalizedDocumentSection] = []
    limitations: list[str] = []
    used = 0
    try:
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row_number > max_rows:
                    limitations.append(f"XLSX sheet {sheet.title} rows truncated at configured cap")
                    break
                values = list(row[:max_columns])
                if len(row) > max_columns:
                    limitations.append(f"XLSX sheet {sheet.title} columns truncated at configured cap")
                cells = [f"{_column_name(index + 1)}{row_number}" for index in range(len(values))]
                row_text = " | ".join(f"{cell}={value}" for cell, value in zip(cells, values) if value is not None)
                if not row_text:
                    continue
                excerpt, truncated = _truncate(row_text, max_excerpt_bytes - used)
                if not excerpt:
                    limitations.append("XLSX excerpts truncated at configured byte cap")
                    break
                sections.append(_section(document_ref, "xlsx", {"sheet": sheet.title, "row": row_number, "cell": cells[0] if cells else None, "cells": cells}, excerpt, limitations=("XLSX excerpts truncated at configured byte cap",) if truncated else (), kind="table_excerpt"))
                used += len(excerpt.encode("utf-8"))
                if truncated:
                    limitations.append("XLSX excerpts truncated at configured byte cap")
                    break
            if used >= max_excerpt_bytes:
                break
    finally:
        workbook.close()
    return tuple(sections), tuple(dict.fromkeys(limitations)), "limited" if limitations else "normalized"


def _limited_document(
    document_ref: str,
    fmt: str,
    data: bytes,
    limitation: str,
    *,
    source_path: str | None = None,
) -> NormalizedDocument:
    """Create an opaque/limited catalog record without inspecting content."""

    return NormalizedDocument(
        document_ref=document_ref,
        format=fmt or "binary",
        content_hash=_sha256(data),
        size_bytes=len(data),
        sections=(),
        limitations=(str(limitation),),
        extraction="limited",
        source_path=source_path,
    )


def _limited_archive_member(
    archive: zipfile.ZipFile,
    info: zipfile.ZipInfo,
    *,
    document_ref: str,
    fmt: str,
    limitation: str,
) -> NormalizedDocument:
    """Create a provenance-complete limited record without buffering a member.

    Large admitted documents remain visible to the Planner, but their raw
    bytes are never materialised in memory merely to build an excerpt catalog.
    Reading to EOF preserves the SHA-256/CRC evidence boundary while keeping
    memory use bounded.
    """

    digest = hashlib.sha256()
    size = 0
    try:
        with archive.open(info, "r") as stream:
            while chunk := stream.read(1024 * 1024):
                digest.update(chunk)
                size += len(chunk)
    except (OSError, RuntimeError, zipfile.BadZipFile, KeyError) as exc:
        # Preserve the digest of bytes that were actually observed.  Never
        # substitute SHA-256(empty) for an unreadable/corrupt member: that
        # would falsely attest to content the catalog did not verify.
        raise _ArchiveMemberReadError(
            f"document member read failed: {exc}",
            observed_hash=digest.hexdigest() if size else None,
            observed_size=size,
        ) from exc
    if size != int(info.file_size):
        raise _ArchiveMemberReadError(
            f"document member size mismatch: {document_ref}",
            observed_hash=digest.hexdigest() if size else None,
            observed_size=size,
        )
    return NormalizedDocument(
        document_ref=document_ref,
        format=fmt,
        content_hash=digest.hexdigest(),
        size_bytes=size,
        sections=(),
        limitations=(str(limitation),),
        extraction="limited",
        source_path=document_ref,
    )


def _read_archive_member(archive: zipfile.ZipFile, info: zipfile.ZipInfo, *, document_ref: str) -> bytes:
    """Read one bounded archive member while retaining CRC/read evidence."""

    digest = hashlib.sha256()
    payload = bytearray()
    size = 0
    try:
        with archive.open(info, "r") as stream:
            while chunk := stream.read(1024 * 1024):
                payload.extend(chunk)
                digest.update(chunk)
                size += len(chunk)
    except (OSError, RuntimeError, zipfile.BadZipFile, KeyError) as exc:
        raise _ArchiveMemberReadError(
            f"document member read failed: {exc}",
            observed_hash=digest.hexdigest() if size else None,
            observed_size=size,
        ) from exc
    if size != int(info.file_size):
        raise _ArchiveMemberReadError(
            f"document member size mismatch: {document_ref}",
            observed_hash=digest.hexdigest() if size else None,
            observed_size=size,
        )
    return bytes(payload)


def _bound_document_sections(
    document: NormalizedDocument,
    max_bytes: int,
    limitation: str,
) -> NormalizedDocument:
    """Trim an extractor result to an aggregate byte budget."""

    if max_bytes < 0:
        max_bytes = 0
    used = 0
    sections: list[NormalizedDocumentSection] = []
    truncated = False
    for section in document.sections:
        if used >= max_bytes:
            truncated = True
            break
        text, was_truncated = _truncate(section.text, max_bytes - used)
        if not text and section.text:
            truncated = True
            break
        limits = (
            tuple(dict.fromkeys((*section.limitations, limitation)))
            if was_truncated
            else section.limitations
        )
        sections.append(
            NormalizedDocumentSection(
                document_ref=section.document_ref,
                format=section.format,
                locator=section.locator,
                text=text,
                content_hash=_sha256(text.encode("utf-8")),
                limitations=limits,
                kind=section.kind,
            )
        )
        used += len(text.encode("utf-8"))
        truncated = truncated or was_truncated
        if was_truncated:
            break
    if len(sections) != len(document.sections):
        truncated = True
    budget_hit = (
        bool(document.sections) and used >= max_bytes
    ) or (
        max_bytes == 0 and document.extraction != "opaque"
    )
    if not truncated and not budget_hit:
        return document
    limitations = tuple(dict.fromkeys((*document.limitations, limitation)))
    return NormalizedDocument(
        document_ref=document.document_ref,
        format=document.format,
        content_hash=document.content_hash,
        size_bytes=document.size_bytes,
        sections=tuple(sections),
        limitations=limitations,
        extraction="limited",
        source_path=document.source_path,
    )


def _normalize_bytes(
    document_ref: str,
    fmt: str,
    data: bytes,
    *,
    max_excerpt_bytes: int,
    max_member_bytes: int | None,
    max_total_bytes: int | None,
    max_entries: int | None,
    max_rows: int,
    max_columns: int,
    max_pdf_pages: int,
    max_pdf_output_bytes: int,
    pdf_timeout_seconds: float,
    pdf_cpu_seconds: int,
    pdf_memory_bytes: int,
    source_path: str | None = None,
) -> NormalizedDocument:
    digest = _sha256(data)
    size = len(data)
    display_path = str(source_path or _source_path_for_ref(document_ref)).replace("\\", "/")
    if fmt not in SUPPORTED_DOCUMENT_FORMATS:
        return NormalizedDocument(document_ref, fmt or "binary", digest, size, (), (f"unsupported document format: {fmt}",), "opaque", display_path)
    if max_member_bytes is not None and size > max_member_bytes:
        return NormalizedDocument(document_ref, fmt, digest, size, (), ("document exceeds extraction byte cap; original retained in data room",), "opaque", display_path)
    try:
        if fmt in {"txt", "md", "rst"}:
            sections, limitations, extraction = _normalize_text(document_ref, fmt, data, max_excerpt_bytes=max_excerpt_bytes)
        elif fmt in {"csv", "tsv"}:
            sections, limitations, extraction = _normalize_delimited(document_ref, fmt, data, max_excerpt_bytes=max_excerpt_bytes, max_rows=max_rows, max_columns=max_columns)
        elif fmt in {"docx", "odt"}:
            sections, limitations, extraction = _normalize_docx_or_odt(document_ref, fmt, data, max_member_bytes=max_member_bytes, max_total_bytes=max_total_bytes, max_excerpt_bytes=max_excerpt_bytes, max_entries=max_entries)
        elif fmt == "pdf":
            sections, limitations, extraction = _normalize_pdf(
                document_ref,
                data,
                max_excerpt_bytes=max_excerpt_bytes,
                max_pages=max_pdf_pages,
                max_output_bytes=max_pdf_output_bytes,
                timeout_seconds=pdf_timeout_seconds,
                cpu_seconds=pdf_cpu_seconds,
                memory_bytes=pdf_memory_bytes,
            )
        elif fmt == "xlsx":
            sections, limitations, extraction = _normalize_xlsx(document_ref, data, max_excerpt_bytes=max_excerpt_bytes, max_rows=max_rows, max_columns=max_columns, max_member_bytes=max_member_bytes, max_total_bytes=max_total_bytes, max_entries=max_entries)
        else:
            sections, limitations, extraction = (), ("document format is not normalizable",), "opaque"
    except UnsafeDocumentArchiveError:
        raise
    except Exception as exc:
        sections, limitations, extraction = (), (f"document extraction failed: {exc.__class__.__name__}: {exc}",), "opaque"
    # Keep whitespace-only parser output out of the normalized catalog.  Such
    # sections are neither useful evidence nor a safe basis for a Planner
    # originalText echo; the raw document remains available by its hash.
    sections = tuple(section for section in sections if section.text.strip())
    if not sections and not limitations:
        limitations = ("document contains no extractable text",)
        extraction = "limited"
    return NormalizedDocument(document_ref, fmt, digest, size, sections, limitations, extraction, display_path)


def normalize_document(
    path: str | Path,
    *,
    document_ref: str | None = None,
    source_path: str | None = None,
    format: str | None = None,
    max_document_bytes: int | None = DEFAULT_MAX_DOCUMENT_BYTES,
    max_excerpt_bytes: int = DEFAULT_MAX_EXCERPT_BYTES,
    max_member_bytes: int | None = DEFAULT_MAX_MEMBER_BYTES,
    max_total_bytes: int | None = DEFAULT_MAX_TOTAL_BYTES,
    max_entries: int | None = DEFAULT_MAX_ARCHIVE_ENTRIES,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_columns: int = DEFAULT_MAX_COLUMNS,
    max_pdf_pages: int = DEFAULT_MAX_PDF_PAGES,
    max_pdf_output_bytes: int = DEFAULT_MAX_PDF_OUTPUT_BYTES,
    pdf_timeout_seconds: float = DEFAULT_PDF_TIMEOUT_SECONDS,
    pdf_cpu_seconds: int = DEFAULT_PDF_CPU_SECONDS,
    pdf_memory_bytes: int = DEFAULT_PDF_MEMORY_BYTES,
) -> NormalizedDocument:
    """Normalize one regular local document without modifying it."""

    source = Path(path).expanduser()
    if not source.is_file() or source.is_symlink():
        raise FileNotFoundError(source)
    ref = str(document_ref or source.name)
    display_path = str(source_path or source.name).replace("\\", "/")
    fmt = _format(source, format)
    digest, size = _file_sha256(source)
    if max_document_bytes is not None and size > max_document_bytes:
        return NormalizedDocument(ref, fmt, digest, size, (), ("document exceeds configured source byte cap; original retained",), "opaque", display_path)
    data = source.read_bytes()
    return _normalize_bytes(
        ref,
        fmt,
        data,
        max_excerpt_bytes=max_excerpt_bytes,
        max_member_bytes=max_member_bytes,
        max_total_bytes=max_total_bytes,
        max_entries=max_entries,
        max_rows=max_rows,
        max_columns=max_columns,
        max_pdf_pages=max_pdf_pages,
        max_pdf_output_bytes=max_pdf_output_bytes,
        pdf_timeout_seconds=pdf_timeout_seconds,
        pdf_cpu_seconds=pdf_cpu_seconds,
        pdf_memory_bytes=pdf_memory_bytes,
        source_path=display_path,
    )


def normalize_document_bytes(
    data: bytes,
    *,
    document_ref: str,
    source_path: str | None = None,
    format: str | None = None,
    max_document_bytes: int | None = DEFAULT_MAX_DOCUMENT_BYTES,
    max_excerpt_bytes: int = DEFAULT_MAX_EXCERPT_BYTES,
    max_member_bytes: int | None = DEFAULT_MAX_MEMBER_BYTES,
    max_total_bytes: int | None = DEFAULT_MAX_TOTAL_BYTES,
    max_entries: int | None = DEFAULT_MAX_ARCHIVE_ENTRIES,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_columns: int = DEFAULT_MAX_COLUMNS,
    max_pdf_pages: int = DEFAULT_MAX_PDF_PAGES,
    max_pdf_output_bytes: int = DEFAULT_MAX_PDF_OUTPUT_BYTES,
    pdf_timeout_seconds: float = DEFAULT_PDF_TIMEOUT_SECONDS,
    pdf_cpu_seconds: int = DEFAULT_PDF_CPU_SECONDS,
    pdf_memory_bytes: int = DEFAULT_PDF_MEMORY_BYTES,
) -> NormalizedDocument:
    """Normalize bytes from a read-only data-room/archive member."""

    if not isinstance(data, (bytes, bytearray, memoryview)):
        raise TypeError("document data must be bytes")
    payload = bytes(data)
    fmt = _format(document_ref, format)
    display_path = str(source_path or _source_path_for_ref(document_ref)).replace("\\", "/")
    if max_document_bytes is not None and len(payload) > max_document_bytes:
        return NormalizedDocument(document_ref, fmt, _sha256(payload), len(payload), (), ("document exceeds configured source byte cap; original retained",), "opaque", display_path)
    return _normalize_bytes(
        document_ref,
        fmt,
        payload,
        max_excerpt_bytes=max_excerpt_bytes,
        max_member_bytes=max_member_bytes,
        max_total_bytes=max_total_bytes,
        max_entries=max_entries,
        max_rows=max_rows,
        max_columns=max_columns,
        max_pdf_pages=max_pdf_pages,
        max_pdf_output_bytes=max_pdf_output_bytes,
        pdf_timeout_seconds=pdf_timeout_seconds,
        pdf_cpu_seconds=pdf_cpu_seconds,
        pdf_memory_bytes=pdf_memory_bytes,
        source_path=display_path,
    )


def ingest_document_catalog(
    source: str | Path,
    *,
    max_documents: int | None = None,
    max_document_bytes: int | None = DEFAULT_MAX_DOCUMENT_BYTES,
    max_excerpt_bytes: int = DEFAULT_MAX_EXCERPT_BYTES,
    max_member_bytes: int | None = DEFAULT_MAX_MEMBER_BYTES,
    max_total_bytes: int | None = DEFAULT_MAX_TOTAL_BYTES,
    max_entries: int | None = DEFAULT_MAX_ARCHIVE_ENTRIES,
    max_rows: int = DEFAULT_MAX_ROWS,
    max_columns: int = DEFAULT_MAX_COLUMNS,
    max_pdf_pages: int = DEFAULT_MAX_PDF_PAGES,
    max_pdf_output_bytes: int = DEFAULT_MAX_PDF_OUTPUT_BYTES,
    pdf_timeout_seconds: float = DEFAULT_PDF_TIMEOUT_SECONDS,
    pdf_cpu_seconds: int = DEFAULT_PDF_CPU_SECONDS,
    pdf_memory_bytes: int = DEFAULT_PDF_MEMORY_BYTES,
    max_parsed_pdfs: int = DEFAULT_MAX_PARSED_PDFS,
    max_pdf_processes: int | None = None,
    max_pdf_total_wall_seconds: float = DEFAULT_MAX_PDF_TOTAL_WALL_SECONDS,
    max_pdf_total_output_bytes: int = DEFAULT_MAX_PDF_TOTAL_OUTPUT_BYTES,
    max_total_normalized_text_bytes: int = DEFAULT_MAX_TOTAL_NORMALIZED_TEXT_BYTES,
    include_opaque_members: bool = True,
    strict_archive_resource_limits: bool = True,
) -> DocumentCatalog:
    """Read a source file or ZIP data-room into a deterministic catalog.

    Only members with supported document extensions are normalised.  Archive
    structure is always validated strictly.  Resource limits may be soft when
    the archive has already passed the native Data Room admission boundary;
    oversized documents then become bounded metadata records instead of
    aborting the mission.
    """

    if max_pdf_processes is not None:
        max_parsed_pdfs = max_pdf_processes
    if not isinstance(include_opaque_members, bool) or not isinstance(strict_archive_resource_limits, bool):
        raise TypeError("document catalog mode flags must be booleans")
    if max_documents is not None and (isinstance(max_documents, bool) or max_documents < 0):
        raise ValueError("max_documents must be non-negative")
    if isinstance(max_parsed_pdfs, bool) or max_parsed_pdfs < 0:
        raise ValueError("max_parsed_pdfs must be non-negative")
    if max_pdf_total_wall_seconds < 0 or max_pdf_total_output_bytes < 0 or max_total_normalized_text_bytes < 0:
        raise ValueError("catalog aggregate budgets must be non-negative")
    path = Path(source).expanduser()
    if path.is_symlink() or not path.is_file():
        raise FileNotFoundError(path)
    if path.suffix.lower() != ".zip":
        document = normalize_document(
            path,
            max_document_bytes=max_document_bytes,
            max_excerpt_bytes=min(max_excerpt_bytes, max_total_normalized_text_bytes),
            max_member_bytes=max_member_bytes,
            max_total_bytes=max_total_bytes,
            max_entries=max_entries,
            max_rows=max_rows,
            max_columns=max_columns,
            max_pdf_pages=max_pdf_pages,
            max_pdf_output_bytes=min(max_pdf_output_bytes, max_pdf_total_output_bytes),
            pdf_timeout_seconds=min(pdf_timeout_seconds, max_pdf_total_wall_seconds) if max_pdf_total_wall_seconds else 0,
            pdf_cpu_seconds=pdf_cpu_seconds,
            pdf_memory_bytes=pdf_memory_bytes,
        )
        document = _bound_document_sections(
            document,
            max_total_normalized_text_bytes,
            "catalog normalized-text aggregate cap reached",
        )
        if document.format == "pdf":
            document = _bound_document_sections(
                document,
                max_pdf_total_output_bytes,
                "catalog PDF output aggregate cap reached",
            )
        return DocumentCatalog((document,))
    infos, _total = _validate_archive(
        path,
        max_member_bytes=max_member_bytes,
        max_total_bytes=max_total_bytes,
        max_entries=max_entries,
        enforce_resource_limits=strict_archive_resource_limits,
    )
    documents: list[NormalizedDocument] = []
    supported_member_count = 0
    admitted_document_bytes = 0
    parsed_pdfs = 0
    pdf_output_bytes = 0
    normalized_text_bytes = 0
    catalog_started = time.monotonic()
    with zipfile.ZipFile(path, "r") as archive:
        for info in sorted(infos, key=lambda value: _safe_member_name(value.filename)):
            if info.is_dir():
                continue
            name = _safe_member_name(info.filename)
            fmt = _format(name)
            admitted_formats = SUPPORTED_DOCUMENT_FORMATS if include_opaque_members else NORMALIZABLE_DOCUMENT_FORMATS
            if fmt not in admitted_formats:
                continue
            supported_member_count += 1
            try:
                exceeds_soft_source_budget = (
                    not strict_archive_resource_limits
                    and (
                        (max_member_bytes is not None and info.file_size > max_member_bytes)
                        or (
                            max_total_bytes is not None
                            and admitted_document_bytes + info.file_size > max_total_bytes
                        )
                    )
                )
                if exceeds_soft_source_budget:
                    document = _limited_archive_member(
                        archive,
                        info,
                        document_ref=name,
                        fmt=fmt,
                        limitation="document exceeds bounded extraction budget; original retained in data room",
                    )
                    documents.append(document)
                    continue
                payload = _read_archive_member(archive, info, document_ref=name)
                admitted_document_bytes += len(payload)
                if max_documents is not None and len(documents) >= max_documents:
                    document = _limited_document(
                        name,
                        fmt,
                        payload,
                        f"document catalog member cap reached ({max_documents}); member remains raw-only",
                        source_path=name,
                    )
                elif fmt == "pdf" and parsed_pdfs >= max_parsed_pdfs:
                    document = _limited_document(
                        name,
                        fmt,
                        payload,
                        f"catalog PDF parse budget exhausted ({max_parsed_pdfs})",
                        source_path=name,
                    )
                elif fmt == "pdf" and (
                    max_pdf_total_wall_seconds <= 0
                    or time.monotonic() - catalog_started >= max_pdf_total_wall_seconds
                    or pdf_output_bytes >= max_pdf_total_output_bytes
                ):
                    document = _limited_document(
                        name,
                        fmt,
                        payload,
                        "catalog aggregate PDF budget exhausted",
                        source_path=name,
                    )
                elif fmt != "pdf" and normalized_text_bytes >= max_total_normalized_text_bytes:
                    document = _limited_document(
                        name,
                        fmt,
                        payload,
                        "catalog normalized-text budget exhausted",
                        source_path=name,
                    )
                else:
                    pdf_timeout = pdf_timeout_seconds
                    pdf_output_cap = max_pdf_output_bytes
                    excerpt_cap = max_excerpt_bytes
                    if fmt == "pdf":
                        parsed_pdfs += 1
                        remaining_wall = max_pdf_total_wall_seconds - (time.monotonic() - catalog_started)
                        pdf_timeout = min(pdf_timeout_seconds, max(0.0, remaining_wall))
                        pdf_output_cap = min(max_pdf_output_bytes, max(0, max_pdf_total_output_bytes - pdf_output_bytes))
                    else:
                        excerpt_cap = min(max_excerpt_bytes, max(0, max_total_normalized_text_bytes - normalized_text_bytes))
                    if (fmt == "pdf" and (pdf_timeout <= 0 or pdf_output_cap <= 0)) or (fmt != "pdf" and excerpt_cap <= 0):
                        document = _limited_document(
                            name,
                            fmt,
                            payload,
                            "catalog aggregate extraction budget exhausted",
                            source_path=name,
                        )
                    else:
                        document = normalize_document_bytes(
                            payload,
                            document_ref=name,
                            source_path=name,
                            format=fmt,
                            max_document_bytes=max_document_bytes,
                            max_excerpt_bytes=excerpt_cap,
                            max_member_bytes=max_member_bytes,
                            max_total_bytes=max_total_bytes,
                            max_entries=max_entries,
                            max_rows=max_rows,
                            max_columns=max_columns,
                            max_pdf_pages=max_pdf_pages,
                            max_pdf_output_bytes=pdf_output_cap,
                            pdf_timeout_seconds=pdf_timeout,
                            pdf_cpu_seconds=pdf_cpu_seconds,
                            pdf_memory_bytes=pdf_memory_bytes,
                        )
                        remaining_text = max_total_normalized_text_bytes - normalized_text_bytes
                        if remaining_text < 0:
                            remaining_text = 0
                        document = _bound_document_sections(
                            document,
                            remaining_text,
                            "catalog normalized-text aggregate cap reached",
                        )
                        if fmt == "pdf":
                            remaining_pdf_output = max_pdf_total_output_bytes - pdf_output_bytes
                            if remaining_pdf_output < 0:
                                remaining_pdf_output = 0
                            document = _bound_document_sections(
                                document,
                                remaining_pdf_output,
                                "catalog PDF output aggregate cap reached",
                            )
            except (OSError, RuntimeError, zipfile.BadZipFile, KeyError) as exc:
                # A CRC/read failure means the member's content identity was
                # never verified.  Fail the archive closed instead of writing
                # a synthetic (especially SHA-256(empty)) hash that a later
                # planner or revision binder could mistake for source bytes.
                raise UnsafeDocumentArchiveError(
                    f"document archive member integrity check failed: {name}"
                ) from exc
            # Account only normalized excerpts.  Limited/opaque records retain
            # raw byte provenance but cannot be used as evidence by the
            # semantic materializer.
            extracted_bytes = sum(len(section.text.encode("utf-8")) for section in document.sections)
            normalized_text_bytes += extracted_bytes
            if fmt == "pdf":
                pdf_output_bytes += extracted_bytes
            documents.append(document)
    limitations: tuple[str, ...] = ()
    if max_documents is not None and supported_member_count > max_documents:
        limitations = (f"document catalog member cap reached ({max_documents}); capped members remain visible as raw-only records",)
    return DocumentCatalog(tuple(documents), limitations=limitations)


# Friendly aliases used by callers that prefer "normalize" terminology.
ingest_document = normalize_document


def ingest_documents(
    sources: Iterable[str | Path | Mapping[str, Any]],
    *,
    max_documents: int | None = None,
    **kwargs: Any,
) -> DocumentCatalog:
    """Normalize several direct files into one deterministic catalog.

    A mapping source may provide ``path``/``uri`` or in-memory ``data`` and a
    stable ``document_ref``/``documentRef``.  This convenience entry point is
    intentionally not used for ZIP data rooms; pass a ZIP path to
    :func:`ingest_document_catalog` so strict archive validation runs.
    """

    if max_documents is not None and (isinstance(max_documents, bool) or max_documents < 0):
        raise ValueError("max_documents must be non-negative")

    def limited_direct_document(source: str | Path | Mapping[str, Any]) -> NormalizedDocument:
        """Hash, but do not extract, sources beyond an explicit cap."""

        if isinstance(source, Mapping):
            ref = source.get("document_ref", source.get("documentRef", source.get("relativePath")))
            fmt = source.get("format")
            data = source.get("data")
            if data is not None:
                if not isinstance(data, (bytes, bytearray, memoryview)):
                    raise TypeError("document mapping data must be bytes")
                payload = bytes(data)
                document_ref = str(ref or "document")
                document_format = _format(document_ref, str(fmt) if fmt is not None else None)
                return _limited_document(
                    document_ref,
                    document_format,
                    payload,
                    f"document catalog member cap reached ({max_documents}); source remains raw-only",
                )
            path_value = source.get("path", source.get("uri"))
            if path_value is None:
                raise ValueError("document mapping needs path or data")
            path = Path(path_value).expanduser()
            document_ref = str(ref) if ref else path.name
            document_format = _format(path, str(fmt) if fmt is not None else None)
        else:
            path = Path(source).expanduser()
            document_ref = path.name
            document_format = _format(path)
        if path.is_symlink() or not path.is_file():
            raise FileNotFoundError(path)
        digest, size = _file_sha256(path)
        return NormalizedDocument(
            document_ref=document_ref,
            format=document_format,
            content_hash=digest,
            size_bytes=size,
            sections=(),
            limitations=(f"document catalog member cap reached ({max_documents}); source remains raw-only",),
            extraction="limited",
            source_path=str(path).replace("\\", "/"),
        )

    documents: list[NormalizedDocument] = []
    limitations: list[str] = []
    for source in sources:
        if max_documents is not None and len(documents) >= max_documents:
            documents.append(limited_direct_document(source))
            continue
        if isinstance(source, Mapping):
            ref = source.get("document_ref", source.get("documentRef", source.get("relativePath")))
            fmt = source.get("format")
            if source.get("data") is not None:
                data = source.get("data")
                if not isinstance(data, (bytes, bytearray, memoryview)):
                    raise TypeError("document mapping data must be bytes")
                documents.append(normalize_document_bytes(bytes(data), document_ref=str(ref or "document"), format=fmt, **kwargs))
            else:
                path = source.get("path", source.get("uri"))
                if path is None:
                    raise ValueError("document mapping needs path or data")
                documents.append(normalize_document(path, document_ref=str(ref) if ref else None, format=fmt, **kwargs))
        else:
            documents.append(normalize_document(source, **kwargs))
    if max_documents is not None and len(documents) > max_documents:
        limitations.append(
            f"document catalog member cap reached ({max_documents}); capped sources remain visible as raw-only records"
        )
    return DocumentCatalog(tuple(documents), limitations=tuple(limitations))


def normalize_documents(source: str | Path | Iterable[str | Path | Mapping[str, Any]], **kwargs: Any) -> DocumentCatalog:
    if isinstance(source, (str, Path)):
        return ingest_document_catalog(source, **kwargs)
    return ingest_documents(source, **kwargs)


extract_document_catalog = ingest_document_catalog
normalize_file = normalize_document


DocumentCatalogEntry = NormalizedDocument
NormalizedSection = NormalizedDocumentSection


__all__ = [
    "DOCUMENT_INGESTION_SCHEMA_VERSION",
    "SUPPORTED_DOCUMENT_FORMATS",
    "DEFAULT_MAX_DOCUMENT_BYTES",
    "DEFAULT_MAX_MEMBER_BYTES",
    "DEFAULT_MAX_TOTAL_BYTES",
    "DEFAULT_MAX_ARCHIVE_ENTRIES",
    "DEFAULT_MAX_PDF_PAGES",
    "DEFAULT_MAX_PDF_OUTPUT_BYTES",
    "DEFAULT_PDF_TIMEOUT_SECONDS",
    "DEFAULT_PDF_CPU_SECONDS",
    "DEFAULT_PDF_MEMORY_BYTES",
    "DEFAULT_MAX_PARSED_PDFS",
    "DEFAULT_MAX_PDF_PROCESSES",
    "DEFAULT_MAX_PDF_TOTAL_WALL_SECONDS",
    "DEFAULT_MAX_PDF_TOTAL_OUTPUT_BYTES",
    "DEFAULT_MAX_TOTAL_NORMALIZED_TEXT_BYTES",
    "DocumentIngestionError",
    "UnsafeDocumentArchiveError",
    "NormalizedDocumentSection",
    "NormalizedDocument",
    "DocumentCatalog",
    "revision_document_ref",
    "revision_qualify_catalog",
    "DocumentCatalogEntry",
    "NormalizedSection",
    "normalize_document",
    "normalize_document_bytes",
    "ingest_document_catalog",
    "normalize_documents",
    "ingest_document",
    "ingest_documents",
    "extract_document_catalog",
    "normalize_file",
]
